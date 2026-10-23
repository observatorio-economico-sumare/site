#!/usr/bin/env python3
"""Baixa sinopses agregadas e filtra linhas municipais, sem versionar arquivos grandes."""
import argparse
import hashlib
import io
import json
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

SOURCES = [
    ("inep_sinopse_2025",
     "https://download.inep.gov.br/dados_abertos/sinopses_estatisticas/sinopse_estatistica_censo_escolar_2025.zip"),
    ("rais_tabelas_2025",
     "https://www.gov.br/trabalho-e-emprego/pt-br/acesso-a-informacao/acoes-e-programas/programas-projetos-acoes-obras-e-atividades/estatisticas-trabalho/rais/rais-2025/rais-2025-tabelas.xlsx"),
]
ALLOWED = {"download.inep.gov.br", "www.gov.br"}
MAX_DOWNLOAD = 80000000
MAX_INNER = 90000000

def norm(value):
    text = unicodedata.normalize("NFKD", str(value))
    return "".join(c for c in text if not unicodedata.combining(c)).lower()

def download(url, path):
    if urlparse(url).hostname not in ALLOWED:
        raise ValueError("Domínio não autorizado")
    h = hashlib.sha256()
    count = 0
    with urlopen(Request(url, headers={"User-Agent":"OES/0.1 research public statistics"}),
                 timeout=55) as response:
        if urlparse(response.geturl()).hostname not in ALLOWED:
            raise ValueError("Redirecionamento não autorizado")
        with path.open("wb") as out:
            while True:
                chunk = response.read(262144)
                if not chunk:
                    break
                count += len(chunk)
                if count > MAX_DOWNLOAD:
                    raise ValueError("Download excedeu limite de 80 MB")
                out.write(chunk)
                h.update(chunk)
    return count, h.hexdigest()

def workbooks(path):
    if path.suffix == ".xlsx":
        yield path.name, path.read_bytes()
        return
    with zipfile.ZipFile(path) as z:
        for item in z.infolist():
            if (item.filename.lower().endswith(".xlsx") and
                    item.file_size < MAX_INNER and not item.is_dir()):
                yield item.filename, z.read(item)

def inspect(out):
    from openpyxl import load_workbook
    report = []
    matches = []
    with tempfile.TemporaryDirectory(prefix="oes_download_") as d:
        for label, url in SOURCES:
            record = {"source":label,"url":url,"status":"error","sha256":"",
                      "bytes":0,"sheets":[],"sumare_rows":0,"error":""}
            path = Path(d) / (label + (".zip" if label.startswith("inep") else ".xlsx"))
            try:
                record["bytes"],record["sha256"] = download(url,path)
                for filename, content in workbooks(path):
                    if len(content)>MAX_INNER:
                        continue
                    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                    for sheet in wb:
                        record["sheets"].append({"workbook":filename,"sheet":sheet.title})
                        if label == "rais_tabelas_2025":
                            preview = []
                            for preview_row in sheet.iter_rows(min_row=1, max_row=14,
                                                               max_col=35, values_only=True):
                                preview.append([str(v)[:120] if v is not None else ""
                                                for v in preview_row])
                            record["sheets"][-1]["header_preview"] = preview
                        if len(record["sheets"])>250:
                            break
                        for row_index, row in enumerate(sheet.iter_rows(values_only=True),start=1):
                            if row_index>100000:
                                break
                            # Matches on municipality name, never personal name or ID.
                            if any("sumare" in norm(v) and len(str(v)) < 100 for v in row if v is not None):
                                values = [str(v)[:150] if v is not None else "" for v in row[:100]]
                                matches.append({"source":label,"workbook":filename,
                                                "sheet":sheet.title,"row":row_index,
                                                "values":values})
                                record["sumare_rows"] += 1
                                if record["sumare_rows"]>=90:
                                    break
                        if record["sumare_rows"]>=90:
                            break
                    wb.close()
                    if record["sumare_rows"]>=90:
                        break
                record["status"] = "extracted_unverified"
            except Exception as exc:
                record["error"] = "%s: %s"%(type(exc).__name__,str(exc)[:300])
            report.append(record)
            print(label,record["status"],"bytes",record["bytes"],
                  "sumare_rows",record["sumare_rows"],"error",record["error"])
    (out/"fontes_planilhas.json").write_text(
        json.dumps({"validated":False,"publication_allowed":False,
                    "sources":report,"candidates":matches},ensure_ascii=False,indent=2)+"\n",
        encoding="utf-8")

if __name__ == "__main__":
    p=argparse.ArgumentParser()
    p.add_argument("--output",type=Path,required=True)
    args=p.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    inspect(args.output)
